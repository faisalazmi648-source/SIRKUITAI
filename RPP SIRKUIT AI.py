import os
import google.generativeai as genai
from docx import Document
from docx.shared import Pt, Inches, RGBColor
import streamlit as st

# ==========================================
# KONFIGURASI API KEY GOOGLE GEMINI
# ==========================================
# Mengambil API key secara aman dari Streamlit Secrets
API_KEY = st.secrets["GEMINI_API_KEY"]
genai.configure(api_key=API_KEY)

# ==========================================
# FUNGSI PERAKIT PROMPT AUTOMATIS
# ==========================================
def buat_prompt_rpp(data_input):
    prompt_template = f"""
Anda adalah seorang guru profesional yang handal dalam membuat rencana kegiatan pembelajaran. buatkan saya kegiatan pembelajaran dalam bentuk documen atau format word. Kegiatan pembelajaran untuk Untuk kelas {data_input.get('kelas')} dengan mata pelajaran {data_input.get('mata_pelajaran')} dengan fokus materi {data_input.get('fokus_materi')}. identitas kegiatan pembelajaran meliputi: nama penulis :{data_input.get('nama_penulis')}; NIP. {data_input.get('nip')}; Nama instansi: {data_input.get('nama_instansi')}; tanggal pembuatan: {data_input.get('tanggal_pembuatan')}; dan semester: {data_input.get('semester')}. kegiatan pembelajaran meliputi  aktifitas guru dan siswanya, alokasi waktu tiap kegiatan, alat dan bahan yang perlu disiapkan, CP, ATP, Profil pelajar pancasila. Kegiatan pembelajaran disesuaikan dengan model pembelajaran di bawah ini: 

Pengembangan model SIRKUIT AI dilandaskan pada kerangka Inkuiri Kolaboratif (Collaborative Inquiry) melalui siklus Assess-Design-Implement-Measure/Reflect/Change. Model ini terdiri dari 7 tahapan yang membentuk akronim S-I-R-K-U-I-T. Berikut adalah penjelasan detail beserta contoh penerapannya yang bisa diadaptasi untuk semua mata pelajaran.

1. S - Selidiki (Orientasi Masalah Nyata) Penjelasan: Guru menyajikan masalah yang benar-benar ada di lingkungan sekitar. Siswa mengamati, membangun rasa ingin tahu, dan merumuskan pertanyaan investigasi.
2. I - Investigasi (Pengumpulan Data) Penjelasan: Secara berkelompok, siswa mengumpulkan data primer (wawancara, observasi fisik, data numerik lapangan, buku cetak).
3. R - Rancang Pertanyaan untuk AI (Prompting) Penjelasan: Setelah punya data, siswa belajar membuat prompt spesifik untuk AI untuk mencari dugaan awal/teori umum.
4. K - Kritis dan Verifikasi (AI Verification Loop) Penjelasan: Siswa membuat Tabel Verifikasi untuk membandingkan klaim AI dengan data hasil Investigasi nyata di lapangan.
5. U - Ubah Menjadi Solusi Penjelasan: Berdasarkan data nyata yang diverifikasi dengan teori AI, siswa merancang solusi logis (karya fisik, algoritma/computational thinking, atau maket).
6. I - Implementasi (Uji Produk) Penjelasan: Kelompok mempresentasikan karya dan diuji oleh kelompok lain sebagai "Tim Penguji" validitas data.
7. T - Tindak Lanjut (Refleksi dan Perbaikan) Penjelasan: Siswa merefleksikan proses belajar dan menyimpulkan bahwa kebenaran mutlak diukur dari realitas lapangan.

Buatkan kegiatan pembelajaran secara sederhana, sistematis, rapi, dan langsung mencakup seluruh 7 tahapan SIRKUIT AI di atas secara rinci untuk mata pelajaran dan materi yang telah ditentukan.
"""
    return prompt_template

# ==========================================
# FUNGSI PEMBUAT FILE DOCX
# ==========================================
def buat_file_docx(teks_rpp, nama_file_output):
    doc = Document()
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    baris_list = teks_rpp.split('\n')
    for baris in baris_list:
        baris_clean = baris.strip()
        if not baris_clean:
            doc.add_paragraph()
            continue
        
        p = doc.add_paragraph()
        parts = baris_clean.split('**')
        for i, part in enumerate(parts):
            if not part:
                continue
            run = p.add_run(part)
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            if i % 2 == 1:
                run.bold = True

    doc.save(nama_file_output)

# ==========================================
# TAMPILAN ANTARMUKA APLIKASI (STREAMLIT)
# ==========================================
st.set_page_config(page_title="Kegiatan Pembelajaran Berbasis SIRKUIT AI", page_icon="📚", layout="centered")

st.title("📚 Kegiatan Pembelajaran (Model SIRKUIT AI)")
st.markdown("Aplikasi berbasis AI untuk menyusun Rencana Pelaksanaan Pembelajaran secara instan dan siap diunduh dalam format Word (`.docx`).")
st.divider()

# Form Input Data di Web
with st.form("form_rpp"):
    st.subheader("✍️ Masukkan Identitas & Materi Pembelajaran")
    
    col1, col2 = st.columns(2)
    with col1:
        nama_penulis = st.text_input("Nama Penulis & Gelar", "Faisal Azmi Bakhtiar, M.Pd.")
        nip = st.text_input("NIP", "19910515 201801 1 002")
        nama_instansi = st.text_input("Nama Instansi / Sekolah", "SD Negeri Negla 04")
    with col2:
        tanggal_pembuatan = st.text_input("Tanggal Pembuatan", "27 Oktober 2026")
        semester = st.selectbox("Semester", ["Ganjil (1)", "Genap (2)"])
        kelas = st.text_input("Kelas", "V (Lima)")

    mata_pelajaran = st.text_input("Mata Pelajaran", "IPAS (Ilmu Pengetahuan Alam dan Sosial)")
    fokus_materi = st.text_input("Fokus Materi / Topik", "Ekosistem dan Keseimbangan Lingkungan Sekolah")

    # Tombol Eksekusi
    submit_button = st.form_submit_button(label="🚀 Buat & Generate RPP Sekarang")

# Proses ketika tombol diklik
if submit_button:
    data_rpp = {
        "nama_penulis": nama_penulis,
        "nip": nip,
        "nama_instansi": nama_instansi,
        "tanggal_pembuatan": tanggal_pembuatan,
        "semester": semester,
        "kelas": kelas,
        "mata_pelajaran": mata_pelajaran,
        "fokus_materi": fokus_materi
    }

    with st.spinner("⏳ Sedang merakit prompt otomatis dan memproses ke mesin AI... Mohon tunggu sebentar."):
        # 1. Buat Prompt
        prompt_final = buat_prompt_rpp(data_rpp)

        # 2. Panggil Gemini API menggunakan google-generativeai stabil
        generation_config = {
            "temperature": 0.4,
        }
        
        # Menggunakan model gemini-3.8-flash dengan system instruction terintegrasi
        model = genai.GenerativeModel(
            model_name="gemini-3.8-flash",
            system_instruction="Anda adalah Pakar Kurikulum dan Guru Penggerak Indonesia yang handal membuat modul ajar.",
            generation_config=generation_config
        )
        
        # Menambahkan request_options timeout di sini agar tidak DeadlineExceeded
        respon = model.generate_content(
            prompt_final,
            request_options={'timeout': 120}
        )
        teks_hasil = respon.text

        # 3. Simpan ke File Word (.docx)
        nama_file = f"RPP_{mata_pelajaran.replace(' ', '_')}_Kelas_{kelas}.docx"
        buat_file_docx(teks_hasil, nama_file)

    st.success("🎉 RPP Berhasil Dibuat!")
    
    # Tombol Unduh File (.docx) di Browser
    with open(nama_file, "rb") as file_dokumen:
        st.download_button(
            label="📥 Unduh Dokumen RPP (Format Word)",
            data=file_dokumen,
            file_name=nama_file,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    # Pratinjau Teks di Layar
    with st.expander("👁️ Lihat Pratinjau Hasil RPP di Layar"):
        st.markdown(teks_hasil)
